"""The workpaper carries the evidence record, run details and case notes."""

import io
from datetime import datetime

import openpyxl
import pandas as pd

from workpaper_export import build_workpaper


def test_workpaper_has_evidence_and_run_details(parse):
    result = parse("1_checking_2025-03.pdf")
    evidence = pd.DataFrame([{"File": "1_checking_2025-03.pdf", "SHA-256": result.source["sha256"]}])
    log = pd.DataFrame([{"time": "t", "event": "Statement added", "detail": "x"}])
    data = build_workpaper(
        result.transactions, result.flags, pd.DataFrame(), datetime.now().astimezone(), None, None,
        evidence, [("App version", "1.1.0"), ("Analyst", "J. Investigator")], log,
    )
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert {"Summary", "Evidence", "Case_Log", "Transactions"} <= set(book.sheetnames)
    summary = {row[0].value: row[1].value for row in book["Summary"].iter_rows(min_row=2)}
    assert summary["App version"] == "1.1.0" and summary["Analyst"] == "J. Investigator"
    assert book["Evidence"]["B2"].value == result.source["sha256"]
    headers = [c.value for c in book["Transactions"][1]]
    assert {"source_file", "source_sha256", "source_page", "source_line"} <= set(headers)
