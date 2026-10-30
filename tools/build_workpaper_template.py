"""
Build the student workpaper template: docs/templates/Forensic_Workpaper_Template.xlsx.

    python tools/build_workpaper_template.py

One workbook, one indexed workpaper per sheet, the way a forensic engagement file is organized:

    Index  the file: case details, the list of workpapers, sign-offs and tickmarks
    A-1    scope: the questions the client asked, and where each is answered
    A-2    document index: every document received, its SHA-256 and its Document Check result
    B-1    bank reconciliation: each statement proved against its own printed figures
    B-2    sources and uses of funds, by month
    C-1    fund tracing: each payment out, matched to where it arrived
    C-2    payee analysis: what each payee was paid against what is supported
    C-3    cash analysis: withdrawals, structuring, and matching deposits
    D-1    red flag log: every flag (from the app or found by hand) and how it was resolved
    D-2    document authenticity: the checks on each questioned document
    E-1    loss calculation
    F-1    findings summary
    R-1    review notes

Blue cells are for input; white cells hold formulas. Each sheet has one example row (from the
fictional Hartwell case in demo/) in gray italics, to delete before starting. Needs openpyxl.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "templates" / "Forensic_Workpaper_Template.xlsx"

# JFVP colors (the app's palette).
NAVY, BLUE, TAUPE, TINT, LINE = "1D2844", "246094", "918E81", "E8EFF6", "C5D4E3"
INPUT_FILL = PatternFill("solid", fgColor="F2F7FC")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=TINT)
EXAMPLE_FONT = Font(name="Calibri", size=10, italic=True, color="7A7A7A")
BODY_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00;[Red](#,##0.00);"-"'
DATE = "mm/dd/yyyy"

ROWS = 30  # blank input rows per table
HEADER_ROW = 8  # tables start here, below the sign-off block and the instructions

# Drop-down lists used by several sheets.
LISTS = {
    "yes_no": ["Yes", "No"],
    "yes_no_partial": ["Yes", "No", "Partial"],
    "doc_result": ["Looks normal", "Worth checking", "May have been altered", "May be AI-generated", "Not checked"],
    "method": ["Check", "Wire", "Zelle / ACH", "Online transfer", "Cash", "Debit card", "Credit card", "Other"],
    "risk": ["High", "Medium", "Low"],
    "status": ["Open", "Resolved: explained", "Resolved: unexplained", "Referred"],
    "basis": ["Traced to subject", "Unsupported payment", "Unexplained", "Supported (exclude)"],
    "recon": ["Reconciled", "Variance", "Needs Review"],
    "check_result": ["Looks normal", "Worth checking", "Note", "Not checked", "N/A"],
}


def _style_cell(cell, *, font=BODY_FONT, fill=None, fmt=None, wrap=True, bold=False, align=None):
    """Font, fill, number format, border and wrapping for one cell."""
    cell.font = Font(name=font.name, size=font.size, italic=font.italic, color=font.color, bold=bold or font.bold, underline=font.underline)
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    cell.border = BOX
    cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)


def _signoff(ws: Worksheet, ref: str, title: str, purpose: str, width: int) -> None:
    """Rows 1 to 6: case name (from the Index), the workpaper's reference and title, sign-offs, and what the sheet is for."""
    ws["A1"] = "=IF(Index!$C$4=\"\",\"[Case name: fill in on the Index]\",Index!$C$4)"
    ws["A1"].font = Font(name="Calibri", size=10, bold=True, color=TAUPE)
    ws["A2"] = f"{ref}  {title}"
    ws["A2"].font = Font(name="Cambria", size=15, bold=True, color=NAVY)
    last = get_column_letter(width)
    prev = get_column_letter(max(width - 1, 2))
    for row, label in ((1, "Prepared by / date"), (2, "Reviewed by / date")):
        ws[f"{prev}{row}"] = label
        ws[f"{prev}{row}"].font = Font(name="Calibri", size=9, color=TAUPE)
        ws[f"{prev}{row}"].alignment = Alignment(horizontal="right")
        _style_cell(ws[f"{last}{row}"], fill=INPUT_FILL)
    ws.merge_cells(start_row=4, start_column=1, end_row=6, end_column=width)
    ws["A4"] = purpose
    ws["A4"].font = Font(name="Calibri", size=10, color="374151")
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A4"].fill = PatternFill("solid", fgColor="F7F7F5")
    ws.row_dimensions[2].height = 22


def _validation(ws: Worksheet, list_name: str, cells: str) -> None:
    """A drop-down list on a range of cells (from LISTS)."""
    dv = DataValidation(type="list", formula1='"' + ",".join(LISTS[list_name]) + '"', allow_blank=True)
    dv.add(cells)
    ws.add_data_validation(dv)


# A column: header, width, kind ("text", "date", "money", "int", "list:<name>", or "=formula with {r}"), total ("sum" or None).
Column = Tuple[str, float, str, Optional[str]]


def _table(ws: Worksheet, columns: Sequence[Column], example: Optional[Sequence] = None, rows: int = ROWS,
           start: int = HEADER_ROW, totals_label: str = "Total") -> Tuple[int, int]:
    """
    A table of input rows under a navy header, with one gray example row first and a totals row at the end.
    Returns (first data row, last data row), the example row included.
    """
    for i, (header, width, _, _) in enumerate(columns, start=1):
        cell = ws.cell(row=start, column=i, value=header)
        _style_cell(cell, font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=HEAD_FILL)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[start].height = 30
    first, last = start + 1, start + rows + (1 if example else 0)
    for r in range(first, last + 1):
        is_example = bool(example) and r == first
        for i, (_, _, kind, _) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=i)
            formula = kind.startswith("=")
            if formula:
                cell.value = kind.format(r=r)
            elif is_example and i - 1 < len(example):
                value = example[i - 1]
                # Example dates are typed as mm/dd/yyyy for readability; store them as real dates.
                cell.value = datetime.strptime(value, "%m/%d/%Y") if kind == "date" and isinstance(value, str) else value
            fmt = MONEY if kind in ("money",) or (formula and "money" in kind) else DATE if kind == "date" else ("0" if kind == "int" else None)
            if formula and "/*money*/" in kind:
                cell.value = kind.replace("/*money*/", "").format(r=r)
                fmt = MONEY
            if formula and "/*int*/" in kind:
                cell.value = kind.replace("/*int*/", "").format(r=r)
                fmt = "0"
            _style_cell(cell, font=EXAMPLE_FONT if is_example else BODY_FONT, fill=None if formula else INPUT_FILL, fmt=fmt,
                        align="right" if fmt in (MONEY, "0") else None)
        for i, (_, _, kind, _) in enumerate(columns, start=1):
            if kind.startswith("list:"):
                _validation(ws, kind[5:], f"{get_column_letter(i)}{first}:{get_column_letter(i)}{last}")
    if any(total for *_, total in columns):
        tr = last + 1
        ws.cell(row=tr, column=1, value=totals_label)
        for i, (_, _, kind, total) in enumerate(columns, start=1):
            cell = ws.cell(row=tr, column=i)
            col = get_column_letter(i)
            if total == "sum":
                cell.value = f"=SUM({col}{first}:{col}{last})"
            elif total == "count":
                cell.value = f"=COUNTA({col}{first}:{col}{last})"
            _style_cell(cell, fill=TOTAL_FILL, bold=True, fmt=MONEY if total == "sum" else None, align="right" if total else None)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    return first, last


def _page_setup(ws: Worksheet, ref: str) -> None:
    """Landscape, fit to one page wide, with the workpaper reference in the footer."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.oddFooter.left.text = f"Workpaper {ref}"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_view.showGridLines = False


def _new(wb: Workbook, ref: str, title: str, purpose: str, width: int) -> Worksheet:
    """A new workpaper sheet named by its reference, with its sign-off block."""
    ws = wb.create_sheet(ref)
    ws.sheet_properties.tabColor = BLUE
    _signoff(ws, ref, title, purpose, width)
    _page_setup(ws, ref)
    return ws


WORKPAPERS: List[Tuple[str, str, str]] = []  # (ref, title, purpose) for the Index


def build() -> Workbook:
    """Every sheet of the template."""
    wb = Workbook()
    index = wb.active
    index.title = "Index"

    def add(ref: str, title: str, purpose: str, width: int) -> Worksheet:
        """Create a workpaper and list it on the Index."""
        WORKPAPERS.append((ref, title, purpose))
        return _new(wb, ref, title, purpose, width)

    # A-1 Scope -----------------------------------------------------------------------------
    ws = add("A-1", "Scope and questions",
             "List each question the client asked (from the engagement letter), in their words. For each, name the workpapers "
             "that answer it and track its status. Note any limits on the work: missing records, documents you could not verify.", 6)
    _table(ws, [("#", 5, "int", None), ("Question asked by the client", 60, "text", None), ("Answered in (WP refs)", 20, "text", None),
                ("Status", 16, "list:status", None), ("Limitations / records still needed", 45, "text", None), ("Notes", 30, "text", None)],
           example=[1, "What money left Mrs. Hartwell's accounts from 3/1 to 5/31/2025, and where did it go?", "B-2, C-1, C-3",
                    "Open", "Casino ATM cash can't be traced from bank records", ""], rows=12)

    # A-2 Document index ---------------------------------------------------------------------
    ws = add("A-2", "Document index",
             "Every document received, when and from whom. Copy each statement's SHA-256 fingerprint from the app's Evidence record "
             "(Download Workpapers tab, or the Evidence sheet of the app's workpaper) so anyone can later confirm it is the same file. "
             "Record the Document Check result for each statement.", 11)
    _table(ws, [("Doc ID", 8, "text", None), ("Description", 38, "text", None), ("Source (who produced it)", 26, "text", None),
                ("How obtained", 18, "text", None), ("Date received", 12, "date", None), ("File name", 30, "text", None),
                ("SHA-256 (from the app)", 30, "text", None), ("Pages", 7, "int", None), ("Period covered", 18, "text", None),
                ("Document Check result", 18, "list:doc_result", None), ("Notes", 30, "text", None)],
           example=["5", "April 2025 checking statement, \"family copy\"", "Dana Pruitt (by email)", "Client", "06/24/2025",
                    "5_checking_2025-04_family_copy.pdf", "(copy from the app)", 1, "04/01-04/30/2025", "May have been altered",
                    "Compare with Doc 2, the bank's copy"])

    # B-1 Bank reconciliation -------------------------------------------------------------------
    ws = add("B-1", "Bank reconciliation (proof of each statement)",
             "Prove each statement before relying on it: beginning balance + deposits - withdrawals should equal the ending balance the "
             "bank printed. Deposits and withdrawals come from the transactions the app read (Transactions tab, or the workpaper). "
             "A difference means a missed or misread line, or an altered statement. Record the app's Bank Reconciliation result too.", 13)
    _table(ws, [("Doc ID", 7, "text", None), ("Bank", 20, "text", None), ("Account (last 4)", 10, "text", None), ("Period", 18, "text", None),
                ("Beginning balance (per statement)", 15, "money", None), ("Deposits (per detail)", 14, "money", "sum"),
                ("Withdrawals (per detail)", 14, "money", "sum"),
                ("Computed ending", 15, "=IF(E{r}=\"\",\"\",E{r}+F{r}-G{r})/*money*/", None),
                ("Ending balance (per statement)", 15, "money", None),
                ("Difference", 13, "=IF(OR(H{r}=\"\",I{r}=\"\"),\"\",ROUND(I{r}-H{r},2))/*money*/", None),
                ("Result", 12, "=IF(J{r}=\"\",\"\",IF(ABS(J{r})<0.005,\"Reconciled\",\"Variance\"))", None),
                ("App's result", 13, "list:recon", None), ("Tickmarks / notes", 30, "text", None)],
           example=["2", "Larkspur Valley Bank", "8701", "04/01-04/30/2025", 16760.41, 52970.67, 50647.92, None, 19083.16, None, None,
                    "Reconciled", "A F R"])

    # B-2 Sources and uses ---------------------------------------------------------------------
    ws = add("B-2", "Sources and uses of funds",
             "One account per block (copy the block for each account). Rename the categories to fit the case. Sources are money in; uses "
             "are money out, grouped so the unusual stands out (payments to the subject, cash, gambling). The check row should be zero: "
             "if it isn't, a transaction is missing or double-counted.", 6)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15
    r = HEADER_ROW
    for i, header in enumerate(["Account: ________________", "Month 1", "Month 2", "Month 3", "Total", "Notes"], start=1):
        _style_cell(ws.cell(row=r, column=i, value=header), font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=HEAD_FILL)
    def line(row: int, label: str, kind: str = "input", formula: Optional[str] = None, bold: bool = False):
        """One row: a label, three months (input or formula) and a total."""
        _style_cell(ws.cell(row=row, column=1, value=label), bold=bold, fill=TOTAL_FILL if kind != "input" else None)
        for c, col in enumerate("BCD", start=2):
            cell = ws.cell(row=row, column=c)
            if formula:
                cell.value = formula.format(c=col)
            _style_cell(cell, fill=INPUT_FILL if kind == "input" else TOTAL_FILL, fmt=MONEY, bold=bold, align="right")
        total = ws.cell(row=row, column=5, value=f"=SUM(B{row}:D{row})" if label != "Opening balance" else "=B{0}".format(row))
        _style_cell(total, fill=TOTAL_FILL, fmt=MONEY, bold=True, align="right")
        _style_cell(ws.cell(row=row, column=6), fill=INPUT_FILL)
    r += 1
    opening = r
    line(r, "Opening balance")
    sources = ["Pension", "Social Security", "CD maturity / transfers from savings", "Other deposits"]
    uses = ["Ordinary living expenses (utilities, groceries, insurance)", "Payments to the subject or related parties",
            "Cash withdrawals (branch and ATM)", "Gambling", "Gift cards", "Fees (overdraft, wire, ATM)", "Other"]
    r += 1
    _style_cell(ws.cell(row=r, column=1, value="Sources"), bold=True)
    first_source = r + 1
    for label in sources:
        r += 1
        line(r, "  " + label)
    total_sources = r + 1
    line(total_sources, "Total sources", kind="formula", formula=f"=SUM({{c}}{first_source}:{{c}}{r})", bold=True)
    r = total_sources + 1
    _style_cell(ws.cell(row=r, column=1, value="Uses"), bold=True)
    first_use = r + 1
    for label in uses:
        r += 1
        line(r, "  " + label)
    total_uses = r + 1
    line(total_uses, "Total uses", kind="formula", formula=f"=SUM({{c}}{first_use}:{{c}}{r})", bold=True)
    ending = total_uses + 1
    line(ending, "Computed ending balance", kind="formula", formula=f"={{c}}{opening}+{{c}}{total_sources}-{{c}}{total_uses}", bold=True)
    ws.cell(row=ending, column=5).value = f"=D{ending}"
    printed = ending + 1
    line(printed, "Ending balance per statement")
    ws.cell(row=printed, column=5).value = f"=D{printed}"
    check = printed + 1
    line(check, "Check (should be 0)", kind="formula", formula=f"=ROUND({{c}}{printed}-{{c}}{ending},2)", bold=True)
    ws.cell(row=check, column=5).value = f"=ROUND(E{printed}-E{ending},2)"
    for c, col in enumerate("CD", start=3):  # each month opens with the month before's ending balance
        ws.cell(row=opening, column=c).value = f"=IF({get_column_letter(c - 1)}{printed}=\"\",{get_column_letter(c - 1)}{ending},{get_column_letter(c - 1)}{printed})"
        ws.cell(row=opening, column=c).fill = TOTAL_FILL
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)

    # C-1 Fund tracing --------------------------------------------------------------------------
    ws = add("C-1", "Fund tracing",
             "Follow each significant payment out of the victim's accounts to where it arrived. One row per payment. Match by amount, date "
             "(usually 0 to 3 business days later) and reference: check endorsements, wire beneficiary details, Zelle confirmation numbers. "
             "Traced = found arriving in a named account. The summary below the table totals what was traced by method.", 15)
    first, last = _table(ws, [
        ("Trace #", 7, "int", None), ("Date out", 11, "date", None), ("From (bank / account)", 20, "text", None),
        ("Description on statement", 34, "text", None), ("Amount out", 13, "money", "sum"), ("Method", 13, "list:method", None),
        ("Recipient", 22, "text", None), ("Receiving bank / account", 22, "text", None), ("Date in", 11, "date", None),
        ("Amount in", 13, "money", "sum"), ("Difference", 12, "=IF(OR(E{r}=\"\",J{r}=\"\"),\"\",ROUND(J{r}-E{r},2))/*money*/", None),
        ("Days", 7, "=IF(OR(B{r}=\"\",I{r}=\"\"),\"\",I{r}-B{r})/*int*/", None),
        ("Traced?", 9, "list:yes_no_partial", None), ("Evidence (doc / WP ref)", 18, "text", None), ("Notes", 28, "text", None)],
        example=[1, "04/15/2025", "Larkspur 8701", "WIRE TRANSFER OUT DP COMPANION CARE LLC REF 88213", 24000, "Wire",
                 "DP Companion Care LLC", "Quillbrook CU 70033390", "04/15/2025", 24000, None, None, "Yes", "Docs 2, 7; Exhibit G",
                 "$18,750 left next day as check 1001 (SUV)"])
    s = last + 3
    _style_cell(ws.cell(row=s, column=4, value="Summary by method"), bold=True, fill=TOTAL_FILL)
    _style_cell(ws.cell(row=s, column=5, value="Amount out"), bold=True, fill=TOTAL_FILL)
    _style_cell(ws.cell(row=s, column=6, value="Traced (Yes)"), bold=True, fill=TOTAL_FILL)
    for i, method in enumerate(LISTS["method"], start=1):
        row = s + i
        _style_cell(ws.cell(row=row, column=4, value=method))
        _style_cell(ws.cell(row=row, column=5, value=f"=SUMIFS($E${first}:$E${last},$F${first}:$F${last},D{row})"), fmt=MONEY, align="right")
        _style_cell(ws.cell(row=row, column=6, value=f"=SUMIFS($J${first}:$J${last},$F${first}:$F${last},D{row},$M${first}:$M${last},\"Yes\")"),
                    fmt=MONEY, align="right")

    # C-2 Payee analysis -------------------------------------------------------------------------
    ws = add("C-2", "Payee analysis: paid versus supported",
             "One row per payee who matters to the case (the app's \"New payee\" findings are a good start). Supported = what an "
             "agreement, invoice or other document justifies. The calculator below works out the most a caregiver could have earned "
             "under an hourly agreement.", 11)
    first, last = _table(ws, [
        ("Payee", 26, "text", None), ("Related to the subject?", 11, "list:yes_no", None), ("First paid", 11, "date", None),
        ("Last paid", 11, "date", None), ("Number of payments", 10, "int", "sum"), ("Total paid", 14, "money", "sum"),
        ("Support (document)", 26, "text", None), ("Supported amount", 14, "money", "sum"),
        ("Unsupported", 14, "=IF(F{r}=\"\",\"\",F{r}-N(H{r}))/*money*/", "sum"), ("WP / doc ref", 14, "text", None), ("Notes", 30, "text", None)],
        example=["DP Companion Care LLC", "Yes", "04/08/2025", "05/09/2025", 5, 29400, "Invoice INV-1001 (Exhibit D)", 576, None,
                 "C-1, Exhibits C, D", "Agreement: $24/hr, 12 hrs/week"], rows=15)
    calc = last + 3
    ws.cell(row=calc, column=1, value="Most that could be earned under an hourly agreement").font = Font(name="Calibri", bold=True, color=NAVY)
    for i, (label, value, fmt) in enumerate([("Hourly rate", None, MONEY), ("Hours per week", None, "0.0"),
                                             ("Weeks in the period", None, "0.0")], start=1):
        _style_cell(ws.cell(row=calc + i, column=1, value=label))
        _style_cell(ws.cell(row=calc + i, column=2, value=value), fill=INPUT_FILL, fmt=fmt, align="right")
    _style_cell(ws.cell(row=calc + 4, column=1, value="Maximum earned"), bold=True, fill=TOTAL_FILL)
    _style_cell(ws.cell(row=calc + 4, column=2, value=f"=N(B{calc + 1})*N(B{calc + 2})*N(B{calc + 3})"), bold=True, fill=TOTAL_FILL, fmt=MONEY, align="right")

    # C-3 Cash analysis -------------------------------------------------------------------------
    ws = add("C-3", "Cash analysis",
             "Every cash withdrawal (branch, ATM, cash advance). The sheet marks amounts from $8,000 to $9,999.99 (just under the $10,000 "
             "reporting threshold) and the days since the previous withdrawal. Then look for matching cash deposits in the subject's "
             "accounts, a few days later and a little less.", 13)
    first, last = _table(ws, [
        ("Date", 11, "date", None), ("Account", 16, "text", None), ("Where (branch / ATM / casino)", 26, "text", None),
        ("Amount", 13, "money", "sum"),
        ("Just under $10,000?", 11, "=IF(D{r}=\"\",\"\",IF(AND(D{r}>=8000,D{r}<10000),\"Yes\",\"No\"))", None),
        ("Days since previous", 10, "=\"\"/*int*/", None),  # filled in below: it refers to the row above
        ("Matching deposit: date", 12, "date", None), ("Deposit account", 18, "text", None), ("Deposit amount", 13, "money", "sum"),
        ("Days later", 9, "=IF(OR(A{r}=\"\",G{r}=\"\"),\"\",G{r}-A{r})/*int*/", None),
        ("Kept by subject?", 10, "list:yes_no_partial", None), ("Evidence", 18, "text", None), ("Notes", 26, "text", None)],
        example=["04/21/2025", "Larkspur 8701", "1st Ave branch, Maple Grove", 9500, None, None, "04/22/2025", "Oxbow 5520", 9000, None,
                 "Yes", "Docs 2, 9", "Deposit $500 less, next day"])
    for r in range(first, last + 1):  # the "previous" row reference: the first row has none
        cell = ws.cell(row=r, column=6)
        cell.value = f"=IF(OR(A{r}=\"\",N(A{r - 1})=0),\"\",A{r}-A{r - 1})" if r > first else ""
    s = last + 3
    for i, (label, formula, fmt) in enumerate([
        ("Total cash out", f"=SUM(D{first}:D{last})", MONEY),
        ("Withdrawals just under $10,000", f"=COUNTIF(E{first}:E{last},\"Yes\")", "0"),
        ("Matched to subject's deposits", f"=SUM(I{first}:I{last})", MONEY),
        ("Share of cash matched", f"=IF(N(B{s})=0,\"\",B{s + 2}/B{s})", "0%")]):
        _style_cell(ws.cell(row=s + i, column=1, value=label), bold=True, fill=TOTAL_FILL)
        _style_cell(ws.cell(row=s + i, column=2, value=formula), bold=True, fill=TOTAL_FILL, fmt=fmt, align="right")

    # D-1 Red flag log --------------------------------------------------------------------------
    ws = add("D-1", "Red flag log",
             "Every red flag: the app's \"Fraud patterns found\" and Document Check warnings, and anything found by hand. A flag is a lead, "
             "not a conclusion. Record what was done to follow it up and whether it was explained.", 11)
    _table(ws, [("#", 5, "int", None), ("Red flag", 34, "text", None), ("Source (app / manual)", 14, "text", None),
                ("Account", 16, "text", None), ("Date(s)", 16, "text", None), ("Amount", 13, "money", "sum"), ("Risk", 9, "list:risk", None),
                ("Follow-up performed", 34, "text", None), ("Result / explanation", 34, "text", None), ("Status", 16, "list:status", None),
                ("WP ref", 10, "text", None)],
           example=[1, "Cash withdrawals kept just under $10,000", "App", "Larkspur 8701", "04/21-05/23/2025", 33380, "High",
                    "Searched subject's statements for cash deposits", "$31,800 deposited by Dana 1-4 days later", "Resolved: unexplained", "C-3"])

    # D-2 Document authenticity ------------------------------------------------------------------
    ws = add("D-2", "Document authenticity",
             "One column per questioned document. Fill in the app's Document Check result for each check, then the checks only a person "
             "can do. A warning is a reason to get the statement directly from the bank; a clean result doesn't prove a document is genuine.", 6)
    checks = [
        "Software that made the PDF", "Saved again after it was created", "Fonts on transaction lines", "Page numbers match the file",
        "Dates inside the statement period", "Transactions in date order", "Descriptions match money in and out",
        "Running balances add up", "Summary agrees with the transactions", "Amounts formatted consistently",
        "Text left by an AI tool", "Dates that exist", "Routing number", "Account number the same throughout",
        "App's overall result",
        "By hand: compared with a copy obtained directly from the bank", "By hand: transactions agree with the other party's records",
        "By hand: claimed transfers appear in the receiving account", "By hand: other observations",
    ]
    for i, header in enumerate(["Check", "Document 1:", "Document 2:", "Document 3:", "Document 4:", "Notes"], start=1):
        _style_cell(ws.cell(row=HEADER_ROW, column=i, value=header), font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=HEAD_FILL)
    ws.column_dimensions["A"].width = 52
    for col in "BCDE":
        ws.column_dimensions[col].width = 20
    ws.column_dimensions["F"].width = 40
    for i, name in enumerate(checks, start=1):
        r = HEADER_ROW + i
        _style_cell(ws.cell(row=r, column=1, value=name), bold=name.startswith("App's"))
        for c in range(2, 7):
            _style_cell(ws.cell(row=r, column=c), fill=INPUT_FILL)
        if not name.startswith(("By hand", "App's")):
            _validation(ws, "check_result", f"B{r}:E{r}")
        elif name.startswith("App's"):
            _validation(ws, "doc_result", f"B{r}:E{r}")
    conclusion = HEADER_ROW + len(checks) + 2
    ws.cell(row=conclusion, column=1, value="Conclusion for each document").font = Font(name="Calibri", bold=True, color=NAVY)
    for c in range(2, 6):
        _style_cell(ws.cell(row=conclusion + 1, column=c), fill=INPUT_FILL)
    ws.row_dimensions[conclusion + 1].height = 60
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)

    # E-1 Loss calculation ------------------------------------------------------------------------
    ws = add("E-1", "Loss calculation",
             "Build the amount up from the workpapers, one line per component, and say what each rests on. Only lines marked Include = Yes "
             "count in the total; the summary splits the total by basis so the reader sees what is traced and what is only unexplained.", 7)
    first, last = _table(ws, [("Component", 40, "text", None), ("Amount", 14, "money", None), ("Basis", 20, "list:basis", None),
                              ("Include?", 9, "list:yes_no", None), ("WP ref", 12, "text", None), ("Calculation / support", 40, "text", None),
                              ("Notes", 30, "text", None)],
                         example=["Payments to DP Companion Care LLC and Dana Pruitt, less invoiced care", 33824, "Unsupported payment", "Yes",
                                  "C-1, C-2", "$34,400 paid - $576 invoiced", ""], rows=15)
    s = last + 2
    _style_cell(ws.cell(row=s, column=1, value="Total included"), bold=True, fill=TOTAL_FILL)
    _style_cell(ws.cell(row=s, column=2, value=f"=SUMIFS(B{first}:B{last},D{first}:D{last},\"Yes\")"), bold=True, fill=TOTAL_FILL, fmt=MONEY, align="right")
    for i, basis in enumerate(LISTS["basis"], start=1):
        _style_cell(ws.cell(row=s + i, column=1, value=f"  of which: {basis}"))
        _style_cell(ws.cell(row=s + i, column=2, value=f"=SUMIFS(B{first}:B{last},D{first}:D{last},\"Yes\",C{first}:C{last},\"{basis}\")"),
                    fmt=MONEY, align="right")

    # F-1 Findings summary --------------------------------------------------------------------------
    ws = add("F-1", "Findings summary",
             "One row per question on A-1. State what the records show, the amount, and the workpapers behind it. Write findings as "
             "\"consistent with\" or \"not consistent with\" the explanations offered; leave intent and guilt to the court. Say what you "
             "could not determine and what record would settle it.", 6)
    _table(ws, [("#", 5, "int", None), ("Question (from A-1)", 34, "text", None), ("Finding", 60, "text", None), ("Amount", 14, "money", None),
                ("Support (WP refs)", 16, "text", None), ("Limitations / not determined", 40, "text", None)],
           example=[4, "Does the Wrenfield statement show the $24,000 was repaid?",
                    "The statement has four signs of an AI-generated document, and the repayments it shows do not appear in Mrs. Hartwell's "
                    "May statement. Not consistent with repayment.", 24000, "A-2, D-2", "Wrenfield Trust Bank not contacted"], rows=10)

    # R-1 Review notes ---------------------------------------------------------------------------
    ws = add("R-1", "Review notes",
             "The reviewer's questions and the preparer's answers. Every note is cleared before the file is final.", 7)
    _table(ws, [("#", 5, "int", None), ("WP ref", 9, "text", None), ("Reviewer's note", 50, "text", None), ("Preparer's response", 50, "text", None),
                ("Cleared by", 12, "text", None), ("Date", 11, "date", None), ("Open?", 8, "list:yes_no", None)], rows=20)

    _index(index)
    return wb


def _index(ws: Worksheet) -> None:
    """The Index: case details, the list of workpapers with links and sign-offs, the color key and the tickmarks."""
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Forensic Accounting Workpapers"
    ws["A1"].font = Font(name="Cambria", size=18, bold=True, color=NAVY)
    ws["A2"] = "Template from the Bank Statement Analyzer · Justice for Fraud Victims Project · Apache 2.0"
    ws["A2"].font = Font(name="Calibri", size=9, color=TAUPE)
    for row, label in ((4, "Case name"), (5, "Client"), (6, "Engagement / matter no."), (7, "Period examined"), (8, "Prepared by"), (9, "Reviewed by")):
        _style_cell(ws.cell(row=row, column=2, value=label), bold=True)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        for c in range(3, 6):
            _style_cell(ws.cell(row=row, column=c), fill=INPUT_FILL)
    start = 11
    for i, header in enumerate(["Ref", "Workpaper", "What it is for", "Prepared (initials / date)", "Reviewed (initials / date)"], start=1):
        _style_cell(ws.cell(row=start, column=i, value=header), font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill=HEAD_FILL)
    for i, (ref, title, purpose) in enumerate(WORKPAPERS, start=1):
        r = start + i
        cell = ws.cell(row=r, column=1, value=ref)
        cell.hyperlink = f"#'{ref}'!A1"
        _style_cell(cell, font=Font(name="Calibri", size=10, color=BLUE, underline="single"), bold=True)
        _style_cell(ws.cell(row=r, column=2, value=title), bold=True)
        _style_cell(ws.cell(row=r, column=3, value=purpose.split(". ")[0].rstrip(".") + "."))
        for c in (4, 5):
            _style_cell(ws.cell(row=r, column=c), fill=INPUT_FILL)
    for col, width in zip("ABCDE", (8, 34, 80, 18, 18)):
        ws.column_dimensions[col].width = width
    r = start + len(WORKPAPERS) + 2
    ws.cell(row=r, column=1, value="Key").font = Font(name="Calibri", bold=True, color=NAVY)
    _style_cell(ws.cell(row=r + 1, column=1), fill=INPUT_FILL)
    ws.cell(row=r + 1, column=2, value="Blue cells: type here")
    _style_cell(ws.cell(row=r + 2, column=1), fill=TOTAL_FILL)
    ws.cell(row=r + 2, column=2, value="Shaded cells: formulas and totals (don't type over them)")
    ex = ws.cell(row=r + 3, column=1, value="Example")
    ex.font = EXAMPLE_FONT
    ws.cell(row=r + 3, column=2, value="Gray italic row on each sheet: an example from the Hartwell case. Delete it before you start.")
    r += 5
    ws.cell(row=r, column=1, value="Tickmarks").font = Font(name="Calibri", bold=True, color=NAVY)
    for i, (mark, meaning) in enumerate([
            ("A", "Agreed to the source document (statement, check image, wire record)"),
            ("F", "Footed: the column adds to the total shown"),
            ("R", "Recomputed: the calculation was redone and agrees"),
            ("T", "Traced to the receiving account"),
            ("N", "No exception noted"),
            ("X", "Exception: see the note")], start=1):
        _style_cell(ws.cell(row=r + i, column=1, value=mark), bold=True, align="center")
        ws.cell(row=r + i, column=2, value=meaning)
    ws.freeze_panes = "A12"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0


def main() -> None:
    """Write the template."""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.properties.creator = "Justice for Fraud Victims Project"
    wb.properties.title = "Forensic accounting workpaper template"
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {len(wb.sheetnames)} sheets ({', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
